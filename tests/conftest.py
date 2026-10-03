"""Shared pytest fixtures for the Resume-Builder test suite."""

from __future__ import annotations

import logging
import tempfile
from pathlib import Path
from typing import Generator, Any

import pytest
from docx import Document
from flask import Flask
from reportlab.pdfgen import canvas

from app import create_app
from config import Config


@pytest.fixture
def tmp_test_dir() -> Generator[Path, None, None]:
    """Provide a temporary directory for test outputs."""
    with tempfile.TemporaryDirectory() as temp_dir:
        yield Path(temp_dir)


@pytest.fixture
def app(tmp_test_dir: Path) -> Generator[Flask, None, None]:
    """Create and configure a Flask application instance for testing."""
    class TestAppConfig(Config):
        TESTING = True
        DEBUG = False
        WTF_CSRF_ENABLED = False
        UPLOAD_FOLDER = tmp_test_dir / "uploads"
        GENERATED_FOLDER = tmp_test_dir / "generated"
        DATABASE_PATH = tmp_test_dir / "resume_builder_test.db"
        LOG_FOLDER = tmp_test_dir / "logs"

    app_instance = create_app(TestAppConfig)
    yield app_instance

    # Cleanup logging handlers after test
    for handler in list(app_instance.logger.handlers):
        handler.close()
        app_instance.logger.removeHandler(handler)
    root_logger = logging.getLogger()
    for handler in list(root_logger.handlers):
        handler.close()
        root_logger.removeHandler(handler)


@pytest.fixture
def client(app: Flask):
    """Flask test client fixture pre-authenticated with test user session."""
    test_client = app.test_client()
    with test_client.session_transaction() as session:
        session["user_name"] = "Test User"
        session["user_email"] = "testuser@example.com"
    return test_client


@pytest.fixture
def unauthenticated_client(app: Flask):
    """Flask test client fixture without active user session."""
    return app.test_client()


@pytest.fixture
def sample_resume_text() -> str:
    """Sample resume text string."""
    return (
        "John Doe\n"
        "Email: john.doe@example.com | Phone: (555) 019-2834\n"
        "LinkedIn: linkedin.com/in/johndoe | GitHub: github.com/johndoe\n\n"
        "PROFESSIONAL SUMMARY\n"
        "Senior Software Engineer with over 6 years of experience building scalable web applications "
        "using Python, Flask, React, and PostgreSQL. Spearheaded API architecture and reduced latency by 35%.\n\n"
        "EXPERIENCE\n"
        "Lead Developer | Tech Corp | 2021 - Present\n"
        "- Engineered microservices using Python and Docker.\n"
        "- Managed a cross-functional team of 8 engineers.\n\n"
        "EDUCATION\n"
        "B.S. in Computer Science | University of Technology | 2020\n\n"
        "SKILLS\n"
        "Python, Flask, Django, SQL, PostgreSQL, Docker, AWS, Git, REST APIs"
    )


@pytest.fixture
def sample_jd_text() -> str:
    """Sample job description text string."""
    return (
        "Senior Software Engineer - Python & Cloud\n\n"
        "Company Overview:\n"
        "We are looking for an experienced Senior Software Engineer to lead backend architecture.\n\n"
        "Responsibilities:\n"
        "- Design and implement scalable RESTful APIs using Python and Flask.\n"
        "- Deploy microservices to AWS cloud infrastructure using Docker.\n"
        "- Collaborate with product managers and junior developers.\n\n"
        "Requirements:\n"
        "- 5+ years of software development experience.\n"
        "- Expert proficiency in Python, SQL, PostgreSQL, Docker, and REST APIs.\n"
        "- Strong understanding of CI/CD pipelines and unit testing."
    )


@pytest.fixture
def sample_docx_file(tmp_test_dir: Path) -> Path:
    """Generate a valid DOCX resume file for testing."""
    docx_path = tmp_test_dir / "sample_resume.docx"
    doc = Document()
    doc.add_heading("Jane Smith", level=1)
    doc.add_paragraph("Email: jane.smith@example.com | Phone: (555) 987-6543")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Software Engineer at ACME Corp. Built python flask applications.")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.S. Computer Engineering, 2019")
    doc.save(docx_path)
    return docx_path


@pytest.fixture
def sample_pdf_file(tmp_test_dir: Path) -> Path:
    """Generate a valid PDF resume file for testing."""
    pdf_path = tmp_test_dir / "sample_resume.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(100, 750, "Jane Smith - PDF Resume")
    c.drawString(100, 730, "Email: jane.smith@example.com | Phone: (555) 987-6543")
    c.drawString(100, 710, "Skills: Python, Flask, SQL, Docker")
    c.save()
    return pdf_path


@pytest.fixture
def sample_resume_data() -> dict[str, Any]:
    """Sample structured resume dictionary."""
    return {
        "personal": {
            "full_name": "Jane Smith",
            "email": "jane.smith@example.com",
            "phone": "555-987-6543",
            "linkedin": "linkedin.com/in/janesmith",
            "github": "github.com/janesmith",
            "portfolio": "janesmith.dev",
            "address": "123 Tech Lane, San Francisco, CA",
        },
        "education": {
            "degree": "B.S. Computer Engineering",
            "college": "Stanford University",
            "graduation_year": "2020",
            "gpa": "3.9",
        },
        "skills": "Python, Flask, JavaScript, SQL, AWS, Docker",
        "experience": {
            "company": "Tech Innovations",
            "role": "Software Engineer",
            "duration": "2020 - Present",
            "responsibilities": "Designed microservices and optimized PostgreSQL database queries.",
        },
        "projects": {
            "project_name": "AI Resume Builder",
            "description": "Full stack Flask application providing ATS analysis and resume generation.",
            "technologies": "Python, Flask, OpenAI, SQLite",
        },
        "certifications": "AWS Certified Solutions Architect",
        "achievements": "Hackathon Winner 2022",
        "languages": "English (Native), Spanish (Intermediate)",
    }
